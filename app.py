import os
import io
import math
import sqlite3
import pandas as pd
import streamlit as st
import plotly.express as px
import plotly.graph_objects as go
from openpyxl.styles import PatternFill
from openpyxl.utils import get_column_letter

# Local Module Imports
from rules import (
    NyayaAuditEngine, ensure_db_schema,
    RuleCT01, RuleCT02, RuleCT03, RuleCT04, RuleCT05,
    RuleCT06, RuleCT07, RuleCT08, RuleCT09, RuleCT10,
    RuleCT11, RuleCT12, RuleCT13, RuleCT14, RuleCT15,
    RuleCT16, RuleCT17, RuleCT18, RuleCT19, RuleCT20, RuleCT21
)
from legal_ai import get_legal_strategy
from parser import parse_and_ingest_text
from generate_cheatsheet import generate_cheatsheet
try:
    from ingest import ingest_bundle
except Exception:
    ingest_bundle = None

DB_PATH = "nyayatrace.db"

# ----------------------------------------------------
# PAGE CONFIGURATION & ENTERPRISE LIGHT THEME
# ----------------------------------------------------
st.set_page_config(
    page_title="NyayaTrace Audit Workbench",
    layout="wide",
    initial_sidebar_state="expanded"
)

# Custom Enterprise CSS Overrides (Clean, Professional, Zero Emojis)
st.markdown("""
<style>
    @import url('https://fonts.googleapis.com/css2?family=Plus+Jakarta+Sans:wght@400;500;600;700;800&family=Inter:wght@400;500;600;700&display=swap');

    /* Global Reset */
    html, body, [class*="css"], .stApp {
        font-family: 'Plus Jakarta Sans', 'Inter', sans-serif !important;
        background-color: #f8fafc !important;
        color: #0f172a !important;
    }

    /* Fix Block Container Top Padding */
    .block-container {
        padding-top: 2.2rem !important;
        padding-bottom: 4rem !important;
        max-width: 1240px !important;
    }

    /* Target Text Elements */
    label[data-testid="stWidgetLabel"], 
    .stMarkdown label {
        color: #0f172a !important;
        font-weight: 600 !important;
    }

    /* Input Boxes Styling */
    div[data-baseweb="input"],
    div[data-baseweb="input"] > div,
    div[data-baseweb="textarea"],
    div[data-baseweb="textarea"] > div,
    div[data-baseweb="select"],
    div[data-baseweb="select"] > div {
        background-color: #ffffff !important;
        border: 1px solid #cbd5e1 !important;
        border-radius: 8px !important;
        color: #0f172a !important;
        box-shadow: none !important;
    }

    div[data-baseweb="input"] input,
    div[data-baseweb="textarea"] textarea,
    div[data-baseweb="select"] input,
    div[data-baseweb="select"] span {
        color: #0f172a !important;
        background-color: #ffffff !important;
        font-size: 0.92rem !important;
        font-family: 'Inter', sans-serif !important;
    }

    div[data-baseweb="input"] input::placeholder,
    div[data-baseweb="textarea"] textarea::placeholder {
        color: #94a3b8 !important;
    }

    /* File Uploader Component Styling */
    section[data-testid="stFileUploader"] {
        background-color: #ffffff !important;
        border: 1.5px dashed #3b82f6 !important;
        border-radius: 10px !important;
        padding: 18px !important;
    }

    section[data-testid="stFileUploader"] * {
        color: #0f172a !important;
    }

    section[data-testid="stFileUploader"] button {
        background-color: #f1f5f9 !important;
        color: #1e293b !important;
        border: 1px solid #cbd5e1 !important;
        border-radius: 6px !important;
        font-weight: 600 !important;
    }

    /* Override Streamlit Buttons */
    div.stButton > button {
        border-radius: 8px !important;
        font-weight: 600 !important;
        font-size: 0.92rem !important;
        padding: 10px 20px !important;
        transition: all 0.2s ease-in-out !important;
    }

    /* Primary Buttons (Royal Blue with WHITE Text) */
    div.stButton > button[kind="primary"],
    button[data-testid="baseButton-primary"] {
        background: #2563eb !important;
        border: none !important;
        box-shadow: 0 2px 6px rgba(37,99,235,0.2) !important;
    }

    div.stButton > button[kind="primary"] *,
    button[data-testid="baseButton-primary"] * {
        color: #ffffff !important;
        fill: #ffffff !important;
    }

    div.stButton > button[kind="primary"]:hover,
    button[data-testid="baseButton-primary"]:hover {
        background: #1d4ed8 !important;
        box-shadow: 0 4px 12px rgba(37,99,235,0.3) !important;
    }

    /* Secondary Buttons */
    div.stButton > button[kind="secondary"],
    button[data-testid="baseButton-secondary"] {
        background-color: #ffffff !important;
        border: 1px solid #cbd5e1 !important;
        box-shadow: 0 1px 2px rgba(0,0,0,0.03) !important;
    }

    div.stButton > button[kind="secondary"] *,
    button[data-testid="baseButton-secondary"] * {
        color: #1e293b !important;
    }

    div.stButton > button[kind="secondary"]:hover,
    button[data-testid="baseButton-secondary"]:hover {
        background-color: #eff6ff !important;
        border-color: #3b82f6 !important;
    }

    div.stButton > button[kind="secondary"]:hover * {
        color: #2563eb !important;
    }

    /* Sidebar Styling */
    section[data-testid="stSidebar"] {
        background-color: #ffffff !important;
        border-right: 1px solid #e2e8f0 !important;
    }

    section[data-testid="stSidebar"] * {
        color: #0f172a !important;
    }

    /* Top Brand Header Banner */
    .minimal-header {
        background: #ffffff;
        border: 1px solid #e2e8f0;
        border-radius: 12px;
        padding: 20px 24px;
        margin-top: 5px;
        margin-bottom: 24px;
        box-shadow: 0 1px 3px rgba(0,0,0,0.03);
        display: flex;
        justify-content: space-between;
        align-items: center;
    }

    .brand-title-text {
        font-size: 1.75rem;
        font-weight: 800;
        color: #1e3a8a;
        letter-spacing: -0.5px;
        margin: 0;
    }

    .brand-sub-text {
        color: #64748b;
        font-size: 0.9rem;
        margin-top: 3px;
    }

    .status-badge-minimal {
        background-color: #eff6ff;
        color: #1d4ed8;
        font-weight: 700;
        font-size: 0.82rem;
        padding: 6px 14px;
        border-radius: 20px;
        border: 1px solid #bfdbfe;
    }

    /* Real Case Cards */
    .case-card-minimal {
        background: #ffffff;
        border: 1px solid #e2e8f0;
        border-radius: 12px;
        padding: 20px;
        height: 100%;
        box-shadow: 0 1px 3px rgba(0,0,0,0.03);
        display: flex;
        flex-direction: column;
        justify-content: space-between;
        transition: all 0.2s ease-in-out;
    }

    .case-card-minimal:hover {
        border-color: #3b82f6;
        box-shadow: 0 6px 16px rgba(37,99,235,0.08);
    }

    .case-title-min {
        font-size: 1.1rem;
        font-weight: 800;
        color: #0f172a;
        margin-bottom: 3px;
    }

    .case-sub-min {
        font-size: 0.85rem;
        color: #64748b;
        margin-bottom: 12px;
    }

    .badge-critical-min {
        background-color: #fef2f2;
        color: #dc2626;
        font-weight: 700;
        font-size: 0.75rem;
        padding: 4px 10px;
        border-radius: 8px;
        display: inline-block;
        margin-right: 4px;
        margin-bottom: 6px;
    }

    .badge-high-min {
        background-color: #fff7ed;
        color: #c2410c;
        font-weight: 700;
        font-size: 0.75rem;
        padding: 4px 10px;
        border-radius: 8px;
        display: inline-block;
        margin-right: 4px;
        margin-bottom: 6px;
    }

    .badge-neutral-min {
        background-color: #f1f5f9;
        color: #334155;
        font-weight: 700;
        font-size: 0.75rem;
        padding: 4px 10px;
        border-radius: 8px;
        display: inline-block;
        margin-right: 4px;
        margin-bottom: 6px;
    }

    /* Intake Section Modern Styling */
    .intake-card-box {
        background: #ffffff;
        border: 1px solid #e2e8f0;
        border-radius: 14px;
        padding: 24px;
        box-shadow: 0 1px 4px rgba(0,0,0,0.04);
        margin-bottom: 24px;
    }

    .intake-card-header {
        font-size: 1.25rem;
        font-weight: 800;
        color: #0f172a;
        margin-bottom: 4px;
        letter-spacing: -0.3px;
    }

    .intake-card-sub {
        font-size: 0.85rem;
        color: #64748b;
        margin-bottom: 16px;
        line-height: 1.4;
    }

    div[data-testid="stFileUploader"] {
        background-color: #f8fafc !important;
        border: 2px dashed #cbd5e1 !important;
        border-radius: 10px !important;
        padding: 10px !important;
        transition: all 0.2s ease-in-out !important;
    }

    div[data-testid="stFileUploader"]:hover {
        border-color: #2563eb !important;
        background-color: #eff6ff !important;
    }

    /* KPI Summary Metric Cards */
    .kpi-card-min {
        background: #ffffff;
        border: 1px solid #e2e8f0;
        border-radius: 12px;
        padding: 18px;
        text-align: center;
        box-shadow: 0 1px 3px rgba(0,0,0,0.03);
    }

    .kpi-val-min {
        font-size: 2.1rem;
        font-weight: 800;
        color: #0f172a;
        margin: 2px 0;
    }

    .kpi-lbl-min {
        font-size: 0.75rem;
        font-weight: 700;
        text-transform: uppercase;
        letter-spacing: 0.8px;
        color: #64748b;
    }

    /* Contradiction Cards */
    .contra-item-min {
        background: #ffffff;
        border: 1px solid #e2e8f0;
        border-radius: 10px;
        padding: 18px;
        margin-bottom: 14px;
        box-shadow: 0 1px 3px rgba(0,0,0,0.03);
    }

    .contra-critical { border-left: 5px solid #dc2626; }
    .contra-high { border-left: 5px solid #ea580c; }
    .contra-medium { border-left: 5px solid #ca8a04; }

    .doc-item-min {
        background: #ffffff;
        border: 1px solid #e2e8f0;
        border-left: 5px solid #2563eb;
        border-radius: 10px;
        padding: 14px;
        margin-bottom: 12px;
        box-shadow: 0 1px 3px rgba(0,0,0,0.03);
    }

    /* Tabs Override */
    button[data-baseweb="tab"] {
        font-weight: 700 !important;
        font-size: 0.92rem !important;
        color: #64748b !important;
    }

    button[data-baseweb="tab"][aria-selected="true"] {
        color: #2563eb !important;
        border-bottom-color: #2563eb !important;
    }
</style>
""", unsafe_allow_html=True)


# ----------------------------------------------------
# DATABASE HELPERS & SAMPLE SEEDING
# ----------------------------------------------------
def get_db_connection():
    if not os.path.exists(DB_PATH) or os.path.getsize(DB_PATH) == 0:
        conn = sqlite3.connect(DB_PATH)
        ensure_db_schema(conn)
        conn.close()
    return sqlite3.connect(DB_PATH)


def run_full_audit(target_bundle: str = None):
    """Runs 21-rule audit engine."""
    conn = get_db_connection()
    ensure_db_schema(conn)
    conn.close()
    
    engine = NyayaAuditEngine(DB_PATH)
    engine.run(bundle_id=target_bundle)


def ensure_sample_bundles_loaded():
    """Seeds sample bundles and processes realistic OCR test document safely."""
    conn = get_db_connection()
    ensure_db_schema(conn)
    cur = conn.cursor()

    # 1. Seed GAMMA if missing
    try:
        cur.execute("SELECT COUNT(*) FROM documents WHERE bundle_id='CASE-2026-GAMMA'")
        if cur.fetchone()[0] == 0 and os.path.exists("sample_data_gamma.sql"):
            with open("sample_data_gamma.sql", "r", encoding="utf-8") as f:
                cur.executescript(f.read())
            conn.commit()
    except Exception as e:
        print(f"Warning seeding GAMMA: {e}")

    # 2. Seed DELTA if missing
    try:
        cur.execute("SELECT COUNT(*) FROM documents WHERE bundle_id='CASE-2026-DELTA'")
        if cur.fetchone()[0] == 0 and os.path.exists("sample_data_delta.sql"):
            with open("sample_data_delta.sql", "r", encoding="utf-8") as f:
                cur.executescript(f.read())
            conn.commit()
    except Exception as e:
        print(f"Warning seeding DELTA: {e}")

    # 3. Ingest REALISTIC TEST if missing
    try:
        cur.execute("SELECT COUNT(*) FROM documents WHERE bundle_id='BUNDLE_REALISTIC_TEST'")
        if cur.fetchone()[0] == 0:
            if os.path.exists("realistic_test_fir_seizure.png") and ingest_bundle is not None:
                ingest_bundle("realistic_test_fir_seizure.png", bundle_id="BUNDLE_REALISTIC_TEST")
    except Exception as e:
        print(f"Warning ingesting BUNDLE_REALISTIC_TEST: {e}")

    conn.close()

    # Run audit to ensure contradictions table is up to date
    try:
        run_full_audit()
    except Exception as e:
        print(f"Audit engine warning: {e}")


@st.cache_data(ttl=3)
def load_contradictions_data(bundle_id: str = None):
    conn = get_db_connection()
    query = """
    SELECT 
        c.rowid AS id,
        c.bundle_id,
        c.rule_id,
        c.severity,
        c.narrative,
        c.statute,
        c.event_ids,
        c.created_at
    FROM contradictions c
    """
    if bundle_id:
        query += " WHERE c.bundle_id = ?"
        df = pd.read_sql_query(query, conn, params=(bundle_id,))
    else:
        df = pd.read_sql_query(query, conn)
        
    conn.close()
    return df


@st.cache_data(ttl=3)
def load_documents_for_bundle(bundle_id: str):
    conn = get_db_connection()
    query = """
    SELECT 
        doc_id, 
        bundle_id, 
        doc_type, 
        COALESCE(file_name, doc_type) AS doc_title, 
        COALESCE(uploaded_at, '2026-09-20 10:00:00') AS doc_date, 
        consent_memo_present, 
        sample_certification_present, 
        search_reasons_recorded, 
        videography_present 
    FROM documents 
    WHERE bundle_id = ? 
    ORDER BY doc_id ASC
    """
    df = pd.read_sql_query(query, conn, params=(bundle_id,))
    conn.close()
    return df


def get_involved_events_and_actors(event_ids_str: str):
    """Retrieves involved events & actors."""
    if not event_ids_str:
        return [], []
    
    conn = get_db_connection()
    cursor = conn.cursor()
    parts = [p.strip() for p in str(event_ids_str).split(",") if p.strip()]
    events_data = []
    actors_data = []

    for p in parts:
        cursor.execute("""
            SELECT event_id, doc_id, event_type, COALESCE(timestamp_iso, timestamp_raw, '2026-09-20 10:00') AS ts, COALESCE(location_raw, 'Trial Location') AS loc, lat, lon
            FROM events WHERE event_id = ? OR CAST(event_id AS TEXT) = ? OR CAST(rowid AS TEXT) = ?
        """, (p, p, p))
        ev_rows = cursor.fetchall()
        for r in ev_rows:
            events_data.append({
                "event_id": r[0], "doc_id": r[1], "event_type": r[2],
                "timestamp": r[3], "location": r[4], "lat": r[5], "lon": r[6]
            })

            cursor.execute("""
                SELECT a.actor_id, COALESCE(a.name_raw, a.name_canonical) AS name, a.role, a.address
                FROM event_actors ea
                JOIN actors a ON ea.actor_id = a.actor_id
                WHERE ea.event_id = ?
            """, (r[0],))
            ac_rows = cursor.fetchall()
            for ac in ac_rows:
                actors_data.append({
                    "actor_id": ac[0], "name": ac[1], "role": ac[2], "address": ac[3], "event_id": r[0]
                })

    conn.close()
    return events_data, actors_data


def generate_excel_bytes(df):
    """Generates styled Excel report."""
    output = io.BytesIO()
    with pd.ExcelWriter(output, engine='openpyxl') as writer:
        cols_order = ['bundle_id', 'rule_id', 'severity', 'narrative', 'statute', 'event_ids', 'created_at']
        df_sheet1 = df[[c for c in cols_order if c in df.columns]]
        df_sheet1.to_excel(writer, sheet_name='All Contradictions', index=False)
        
        pivot_df = pd.pivot_table(df, index='rule_id', columns='severity', values='id', aggfunc='count', fill_value=0).reset_index()
        pivot_df.to_excel(writer, sheet_name='Summary', index=False)
        
        bundle_df = df.groupby('bundle_id').size().reset_index(name='total_contradictions')
        bundle_df.to_excel(writer, sheet_name='By Bundle', index=False)
        
        wb = writer.book
        ws_all = wb['All Contradictions']
        red_fill = PatternFill(start_color="FF9999", end_color="FF9999", fill_type="solid")
        orange_fill = PatternFill(start_color="FFCC99", end_color="FFCC99", fill_type="solid")
        yellow_fill = PatternFill(start_color="FFFF99", end_color="FFFF99", fill_type="solid")
        ws_all.freeze_panes = 'A2'
        
        for row in range(2, ws_all.max_row + 1):
            severity_cell = ws_all.cell(row=row, column=3)
            val = str(severity_cell.value).upper() if severity_cell.value else ""
            if val == 'CRITICAL':
                severity_cell.fill = red_fill
            elif val == 'HIGH':
                severity_cell.fill = orange_fill
            elif val == 'MEDIUM':
                severity_cell.fill = yellow_fill

        for sheet in wb.worksheets:
            for col in sheet.columns:
                max_len = max(len(str(cell.value or '')) for cell in col)
                col_letter = get_column_letter(col[0].column)
                sheet.column_dimensions[col_letter].width = min(max(max_len + 3, 12), 70)
                
    output.seek(0)
    return output.getvalue()


# Ensure database has initial sample bundles pre-populated safely
ensure_sample_bundles_loaded()

# ----------------------------------------------------
# MAIN APP HEADER & BRANDING
# ----------------------------------------------------
st.markdown("""
<div class="minimal-header">
    <div>
        <div class="brand-title-text">NyayaTrace Audit Workbench</div>
        <div class="brand-sub-text">Trial Court Evidence & Contradiction Audit Engine (BNSS / NDPS / IPC)</div>
    </div>
    <div class="status-badge-minimal">
        21 Legal Rules Active
    </div>
</div>
""", unsafe_allow_html=True)


# ----------------------------------------------------
# ZERO-DATA / LANDING STATE & REAL CASE EXAMPLES
# ----------------------------------------------------
if "active_bundle" not in st.session_state:
    st.session_state["active_bundle"] = None

if not st.session_state["active_bundle"]:
    up_col1, up_col2 = st.columns([1.1, 1], gap="large")

    with up_col1:
        st.markdown("""
        <div class="intake-card-box">
            <div class="intake-card-header">1. Upload Scanned Trial Bundle</div>
            <div class="intake-card-sub">Drag & drop scanned PDF, PNG, JPG, or TIFF trial court bundle for automated OCR intake and rule auditing.</div>
        """, unsafe_allow_html=True)
        
        uploaded_file = st.file_uploader("Drop trial bundle document here:", type=["pdf", "png", "jpg", "jpeg", "tiff"])
        custom_bundle_id = st.text_input("Assign Case Bundle ID:", "CASE-2026-UPLOADED")
        
        st.markdown("<div style='margin-top:12px;'></div>", unsafe_allow_html=True)
        btn_upload_click = st.button("Process OCR & Audit Uploaded Bundle", type="primary", use_container_width=True, key="btn_audit_file")
        st.markdown("</div>", unsafe_allow_html=True)

        if uploaded_file is not None and btn_upload_click:
            with st.spinner("Processing OCR & Evaluating 21 Audit Rules..."):
                temp_path = os.path.join("debug_ocr", uploaded_file.name)
                os.makedirs("debug_ocr", exist_ok=True)
                with open(temp_path, "wb") as f:
                    f.write(uploaded_file.getbuffer())
                
                # Wipe previous entries for custom_bundle_id to prevent stale data overlap
                conn = get_db_connection()
                cur = conn.cursor()
                cur.execute("DELETE FROM contradictions WHERE bundle_id = ?", (custom_bundle_id,))
                cur.execute("DELETE FROM documents WHERE bundle_id = ?", (custom_bundle_id,))
                conn.commit()
                conn.close()

                if ingest_bundle is not None:
                    ingest_bundle(temp_path, bundle_id=custom_bundle_id)
                run_full_audit(target_bundle=custom_bundle_id)
                
                st.session_state["active_bundle"] = custom_bundle_id
                st.cache_data.clear()
                st.success(f"Audit completed for `{custom_bundle_id}`!")
                st.rerun()

    with up_col2:
        st.markdown("""
        <div class="intake-card-box">
            <div class="intake-card-header">2. Or Paste Raw Document Text</div>
            <div class="intake-card-sub">Paste raw text from FIRs, Seizure Memos, GD Entries, or Witness Statements for instant audit rule analysis.</div>
        """, unsafe_allow_html=True)
        
        pasted_text = st.text_area("Paste FIR or Seizure Memo Text:", height=152, placeholder="FIR Registration: 2026-09-20 10:00:00...\nSeizure Event Timestamp: 2026-09-20 09:15:00...\nLocation of Seizure: Janpath Market, New Delhi...\nConsent Memo Present: 0...")
        
        st.markdown("<div style='margin-top:12px;'></div>", unsafe_allow_html=True)
        if st.button("Audit Pasted Text", type="primary", use_container_width=True, key="btn_audit_text"):
            if pasted_text.strip():
                with st.spinner("Parsing text & evaluating rules..."):
                    parse_and_ingest_text(pasted_text, bundle_id=custom_bundle_id)
                    run_full_audit(target_bundle=custom_bundle_id)
                    st.session_state["active_bundle"] = custom_bundle_id
                    st.cache_data.clear()
                    st.rerun()
        st.markdown("</div>", unsafe_allow_html=True)

    st.markdown("<br>", unsafe_allow_html=True)
    st.markdown("### Real Case Bundles (Select to Audit)")

    c1, c2, c3 = st.columns(3)

    with c1:
        st.markdown("""
        <div class="case-card-minimal">
            <div>
                <div class="case-title-min">State vs. Rajesh Sharma</div>
                <div class="case-sub-min">Karol Bagh NDPS Recovery & Timestamp Paradox</div>
                <div>
                    <span class="badge-neutral-min">6 Contradictions</span>
                    <span class="badge-critical-min">CRITICAL: CT-01</span>
                    <span class="badge-high-min">HIGH: CT-12</span>
                </div>
                <p style="font-size:0.84rem; color:#64748b; margin-top:10px; line-height:1.4;">
                    Scanned police bundle (PS Karol Bagh to Ajmal Khan Market, 0.1 km) with FIR registered after seizure.
                </p>
            </div>
        </div>
        """, unsafe_allow_html=True)
        
        st.markdown("<div style='margin-top:10px;'></div>", unsafe_allow_html=True)
        if st.button("Load & Audit Bundle", key="btn_real_1", use_container_width=True, type="primary"):
            with st.spinner("Loading State vs. Rajesh Sharma..."):
                if os.path.exists("realistic_test_fir_seizure.png") and ingest_bundle is not None:
                    ingest_bundle("realistic_test_fir_seizure.png", bundle_id="BUNDLE_REALISTIC_TEST")
                run_full_audit("BUNDLE_REALISTIC_TEST")
                st.session_state["active_bundle"] = "BUNDLE_REALISTIC_TEST"
                st.cache_data.clear()
                st.rerun()

    with c2:
        st.markdown("""
        <div class="case-card-minimal">
            <div>
                <div class="case-title-min">State vs. Amit Dev</div>
                <div class="case-sub-min">Connaught Place Commercial NDPS Bundle</div>
                <div>
                    <span class="badge-neutral-min">6 Contradictions</span>
                    <span class="badge-critical-min">CRITICAL: CT-01</span>
                    <span class="badge-high-min">HIGH: CT-13</span>
                </div>
                <p style="font-size:0.84rem; color:#64748b; margin-top:10px; line-height:1.4;">
                    New Delhi trial bundle (PS Connaught Place to Janpath Market, 0.8 km) with missing Sec 52A sample cert.
                </p>
            </div>
        </div>
        """, unsafe_allow_html=True)

        st.markdown("<div style='margin-top:10px;'></div>", unsafe_allow_html=True)
        if st.button("Load & Audit Bundle", key="btn_real_2", use_container_width=True, type="primary"):
            with st.spinner("Loading State vs. Amit Dev..."):
                if os.path.exists("new_trial_court_bundle_2026.pdf") and ingest_bundle is not None:
                    ingest_bundle("new_trial_court_bundle_2026.pdf", bundle_id="CASE-2026-DELHI-NDPS")
                run_full_audit("CASE-2026-DELHI-NDPS")
                st.session_state["active_bundle"] = "CASE-2026-DELHI-NDPS"
                st.cache_data.clear()
                st.rerun()

    with c3:
        st.markdown("""
        <div class="case-card-minimal">
            <div>
                <div class="case-title-min">State vs. Vikram Singh</div>
                <div class="case-sub-min">Interstate Teleportation & Bilocation Case</div>
                <div>
                    <span class="badge-neutral-min">7 Contradictions</span>
                    <span class="badge-critical-min">CRITICAL: CT-02</span>
                    <span class="badge-high-min">HIGH: CT-05</span>
                </div>
                <p style="font-size:0.84rem; color:#64748b; margin-top:10px; line-height:1.4;">
                    Multi-document trial bundle with IO traveling 394 km from Karol Bagh to Kanpur in 45 minutes.
                </p>
            </div>
        </div>
        """, unsafe_allow_html=True)

        st.markdown("<div style='margin-top:10px;'></div>", unsafe_allow_html=True)
        if st.button("Load & Audit Bundle", key="btn_real_3", use_container_width=True, type="primary"):
            with st.spinner("Loading State vs. Vikram Singh..."):
                run_full_audit("CASE-2026-DELTA")
                st.session_state["active_bundle"] = "CASE-2026-DELTA"
                st.cache_data.clear()
                st.rerun()

    st.stop() # Stop rendering full dashboard until bundle is loaded


# ----------------------------------------------------
# AUDITED WORKBENCH DASHBOARD STATE (AFTER BUNDLE SELECT)
# ----------------------------------------------------
selected_bundle = st.session_state["active_bundle"]
df_all = load_contradictions_data()
df_bundle = df_all[df_all['bundle_id'] == selected_bundle]

# Sidebar Controls
st.sidebar.markdown(f"### NyayaTrace Controls")
st.sidebar.markdown(f"**Active Case Bundle:** `{selected_bundle}`")

if st.sidebar.button("Upload New Case Bundle", use_container_width=True, type="primary"):
    st.session_state["active_bundle"] = None
    st.rerun()

st.sidebar.markdown("---")

if st.sidebar.button("Re-Run Audit Engine", use_container_width=True):
    with st.spinner("Evaluating 21 legal rules..."):
        run_full_audit(target_bundle=selected_bundle)
        st.cache_data.clear()
        st.sidebar.success("Audit re-run completed!")
        st.rerun()

excel_bytes = generate_excel_bytes(df_all)
st.sidebar.download_button(
    label="Export Excel Audit Report",
    data=excel_bytes,
    file_name="NyayaTrace_Audit_Report.xlsx",
    mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    use_container_width=True
)

st.sidebar.markdown("---")
st.sidebar.subheader("Filters")

# Bundle Switcher
all_bundles = sorted(list(df_all['bundle_id'].unique()))
if selected_bundle in all_bundles:
    bundle_idx = all_bundles.index(selected_bundle)
else:
    bundle_idx = 0

switch_bundle = st.sidebar.selectbox("Switch Active Bundle", all_bundles, index=bundle_idx)
if switch_bundle != selected_bundle:
    st.session_state["active_bundle"] = switch_bundle
    st.rerun()

selected_severity = st.sidebar.selectbox("Filter Severity", ["All Severities", "CRITICAL", "HIGH", "MEDIUM"])

# Filtered data
df_filtered = df_bundle.copy()
if selected_severity != "All Severities":
    df_filtered = df_filtered[df_filtered['severity'] == selected_severity]


# ----------------------------------------------------
# METRIC CARDS ROW
# ----------------------------------------------------
total_cnt = len(df_filtered)
critical_cnt = len(df_filtered[df_filtered['severity'] == 'CRITICAL'])
high_cnt = len(df_filtered[df_filtered['severity'] == 'HIGH'])
medium_cnt = len(df_filtered[df_filtered['severity'] == 'MEDIUM'])

m1, m2, m3, m4 = st.columns(4)
m1.markdown(f'<div class="kpi-card-min"><div class="kpi-lbl-min">Total Violations</div><div class="kpi-val-min">{total_cnt}</div></div>', unsafe_allow_html=True)
m2.markdown(f'<div class="kpi-card-min"><div class="kpi-lbl-min">Critical Severities</div><div class="kpi-val-min" style="color:#dc2626;">{critical_cnt}</div></div>', unsafe_allow_html=True)
m3.markdown(f'<div class="kpi-card-min"><div class="kpi-lbl-min">High Severities</div><div class="kpi-val-min" style="color:#ea580c;">{high_cnt}</div></div>', unsafe_allow_html=True)
m4.markdown(f'<div class="kpi-card-min"><div class="kpi-lbl-min">Medium Severities</div><div class="kpi-val-min" style="color:#ca8a04;">{medium_cnt}</div></div>', unsafe_allow_html=True)

st.markdown("<br>", unsafe_allow_html=True)

# ----------------------------------------------------
# TABS INTERFACE
# ----------------------------------------------------
tab_workbench, tab_analytics, tab_map, tab_cheatsheet = st.tabs([
    "Evidence & Contradictions Workbench",
    "Analytics Overview", 
    "GIS Travel Map",
    "Cross-Exam Cheat Sheet"
])

# ----------------------------------------------------
# TAB 1: SPLIT-SCREEN WORKBENCH
# ----------------------------------------------------
with tab_workbench:
    left_col, right_col = st.columns([1, 1.25])
    
    with left_col:
        st.subheader("Audited Trial Bundle Documents")
        docs_df = load_documents_for_bundle(selected_bundle)
        
        if docs_df.empty:
            st.info(f"No documents found for `{selected_bundle}`.")
        else:
            st.write(f"Documents Parsed: **{len(docs_df)}**")
            for idx, doc in docs_df.iterrows():
                st.markdown(f"""
                <div class="doc-item-min">
                    <div style="font-weight:700; font-size:0.98rem; color:#1e3a8a;">Document: {doc['doc_title']} ({doc['doc_type']})</div>
                    <div style="font-size:0.82rem; color:#64748b; margin-top:2px;">
                        <b>Date:</b> {doc['doc_date']}
                    </div>
                    <div style="font-size:0.8rem; margin-top:6px; color:#1e293b;">
                        Sec 50 Consent: {'Present' if doc['consent_memo_present']==1 else '0 (Violation)'} | 
                        Sec 52A Cert: {'Present' if doc['sample_certification_present']==1 else '0 (Violation)'} | 
                        Videography: {'Present' if doc['videography_present']==1 else '0 (Violation)'}
                    </div>
                </div>
                """, unsafe_allow_html=True)

    with right_col:
        st.subheader("Detected Evidence Contradictions")
        if df_filtered.empty:
            st.success("No contradictions match the selected severity filter.")
        else:
            st.write(f"Violations Found: **{len(df_filtered)}** (Deterministic Engine)")
            for idx, row in df_filtered.iterrows():
                sev = row['severity']
                card_class = "contra-critical" if sev == "CRITICAL" else ("contra-high" if sev == "HIGH" else "contra-medium")
                
                st.markdown(f"""
                <div class="contra-item-min {card_class}">
                    <div style="display:flex; justify-content:space-between; align-items:center;">
                        <span style="font-weight:700; font-size:0.98rem; color:#0f172a;">[{row['rule_id']}] {sev}</span>
                        <span style="font-size:0.78rem; background:#f1f5f9; padding:2px 8px; border-radius:4px; color:#334155; font-weight:700;">Statute: {row['statute']}</span>
                    </div>
                    <div style="font-size:0.9rem; margin-top:6px; line-height:1.4; color:#1e293b;">
                        {row['narrative']}
                    </div>
                </div>
                """, unsafe_allow_html=True)
                
                with st.expander(f"Drilldown Events & Actors ({row['event_ids']})"):
                    events, actors = get_involved_events_and_actors(row['event_ids'])
                    st.markdown("**Involved Events:**")
                    if events:
                        for ev in events:
                            st.markdown(f"- **Event #{ev['event_id']}**: {ev['event_type']} at `{ev['location']}` ({ev['timestamp']})")
                    else:
                        st.write(f"Referenced Events: `{row['event_ids']}`")

                    st.markdown("**Involved Actors:**")
                    if actors:
                        for ac in actors:
                            st.markdown(f"""
                            <div style="background:#f8fafc; border:1px solid #e2e8f0; padding:8px 12px; border-radius:6px; margin-top:6px; font-size:0.85rem;">
                                <b>{ac['name']}</b> ({ac['role']})<br>
                                <i>Address/Station:</i> {ac['address'] or 'Recorded in trial memo'}
                            </div>
                            """, unsafe_allow_html=True)


# ----------------------------------------------------
# TAB 2: ANALYTICS DASHBOARD
# ----------------------------------------------------
with tab_analytics:
    st.subheader("Analytics Overview")
    col_l, col_r = st.columns(2)
    with col_l:
        sev_counts = df_bundle['severity'].value_counts().reset_index()
        sev_counts.columns = ['severity', 'count']
        color_map = {'CRITICAL': '#dc2626', 'HIGH': '#ea580c', 'MEDIUM': '#ca8a04'}
        fig_donut = px.pie(sev_counts, names='severity', values='count', hole=0.5, color='severity', color_discrete_map=color_map, title=f"Severity Breakdown for {selected_bundle}")
        fig_donut.update_layout(template="plotly_white")
        st.plotly_chart(fig_donut, use_container_width=True)
        
    with col_r:
        rule_counts = df_bundle['rule_id'].value_counts().reset_index()
        rule_counts.columns = ['rule_id', 'count']
        fig_bar = px.bar(rule_counts, x='count', y='rule_id', orientation='h', color='count', color_continuous_scale="Reds", title=f"Top Violated Rules for {selected_bundle}")
        fig_bar.update_layout(template="plotly_white")
        st.plotly_chart(fig_bar, use_container_width=True)


# ----------------------------------------------------
# TAB 3: GIS MAP
# ----------------------------------------------------
with tab_map:
    st.subheader(f"Physical Impossibility & Bilocation Map (`{selected_bundle}`)")
    
    conn = get_db_connection()
    events_map_df = pd.read_sql_query("""
        SELECT e.event_id, e.event_type, COALESCE(e.timestamp_iso, e.timestamp_raw, '2026-09-20 10:00') AS ts, 
               COALESCE(e.location_raw, 'Trial Location') AS loc, e.lat, e.lon
        FROM events e
        JOIN documents d ON e.doc_id = d.doc_id
        WHERE d.bundle_id = ? AND e.lat IS NOT NULL AND e.lon IS NOT NULL
        ORDER BY e.timestamp_iso ASC
    """, conn, params=(selected_bundle,))
    conn.close()

    if len(events_map_df) >= 2:
        ev1 = events_map_df.iloc[0]
        ev2 = events_map_df.iloc[1]
        
        lat1, lon1 = float(ev1['lat']), float(ev1['lon'])
        lat2, lon2 = float(ev2['lat']), float(ev2['lon'])
        
        # Calculate Haversine distance
        R = 6371.0
        dlat = math.radians(lat2 - lat1)
        dlon = math.radians(lon2 - lon1)
        a = math.sin(dlat / 2.0)**2 + math.cos(math.radians(lat1)) * math.cos(math.radians(lat2)) * math.sin(dlon / 2.0)**2
        c_angle = 2.0 * math.atan2(math.sqrt(a), math.sqrt(max(0.0, 1.0 - a)))
        dist_km = R * c_angle

        # Calculate time gap
        ts1 = pd.to_datetime(ev1['ts'], errors='coerce')
        ts2 = pd.to_datetime(ev2['ts'], errors='coerce')
        
        if pd.notnull(ts1) and pd.notnull(ts2):
            time_gap_mins = abs((ts2 - ts1).total_seconds()) / 60.0
        else:
            time_gap_mins = 45.0
            
        time_gap_mins = max(time_gap_mins, 1.0)
        speed_kmh = dist_km / (time_gap_mins / 60.0)
        
        # Determine dynamic zoom level based on distance
        if dist_km < 0.5:
            dynamic_zoom = 13
        elif dist_km < 2.0:
            dynamic_zoom = 12
        elif dist_km < 20.0:
            dynamic_zoom = 10
        elif dist_km < 100.0:
            dynamic_zoom = 8
        else:
            dynamic_zoom = 5

        center_lat = (lat1 + lat2) / 2.0
        center_lon = (lon1 + lon2) / 2.0

        map_data = pd.DataFrame([
            {"Location": f"Event A: {ev1['loc']}", "lat": lat1, "lon": lon1, "Time": str(ev1['ts']), "Size": 25, "Event": "Event A"},
            {"Location": f"Event B: {ev2['loc']}", "lat": lat2, "lon": lon2, "Time": str(ev2['ts']), "Size": 25, "Event": "Event B"}
        ])
        
        try:
            if hasattr(px, "scatter_map"):
                fig_map = px.scatter_map(
                    map_data, lat="lat", lon="lon",
                    hover_name="Location", hover_data=["Time"],
                    size="Size", zoom=dynamic_zoom, height=440,
                    center=dict(lat=center_lat, lon=center_lon),
                    title=f"Travel Route Analysis — {selected_bundle}"
                )
                fig_map.add_trace(go.Scattermap(
                    mode="lines", lon=[lon1, lon2], lat=[lat1, lat2],
                    line=dict(width=4, color="#dc2626" if speed_kmh > 150 else "#2563eb"),
                    name=f"Path ({dist_km:.1f} km)"
                ))
                fig_map.update_layout(
                    map_style="open-street-map",
                    template="plotly_white",
                    margin=dict(l=0, r=0, t=40, b=0)
                )
                st.plotly_chart(fig_map, use_container_width=True)
            elif hasattr(px, "scatter_mapbox"):
                fig_map = px.scatter_mapbox(
                    map_data, lat="lat", lon="lon",
                    hover_name="Location", hover_data=["Time"],
                    size="Size", zoom=dynamic_zoom, height=440,
                    center=dict(lat=center_lat, lon=center_lon),
                    title=f"Travel Route Analysis — {selected_bundle}"
                )
                fig_map.add_trace(go.Scattermapbox(
                    mode="lines", lon=[lon1, lon2], lat=[lat1, lat2],
                    line=dict(width=4, color="#dc2626" if speed_kmh > 150 else "#2563eb"),
                    name=f"Path ({dist_km:.1f} km)"
                ))
                fig_map.update_layout(
                    mapbox_style="open-street-map",
                    template="plotly_white",
                    margin=dict(l=0, r=0, t=40, b=0)
                )
                st.plotly_chart(fig_map, use_container_width=True)
            else:
                raise ValueError("Use Plotly Fallback")
        except Exception:
            # Smart fallback with padded coordinate bounds so labels never overlap
            pad_lat = max(0.01, abs(lat1 - lat2) * 0.5)
            pad_lon = max(0.01, abs(lon1 - lon2) * 0.5)
            
            fig_fallback = px.scatter(
                map_data, x="lon", y="lat", text="Location", size="Size", color="Event",
                color_discrete_map={"Event A": "#2563eb", "Event B": "#dc2626"},
                title=f"Travel Route Analysis — {selected_bundle} ({dist_km:.1f} km)"
            )
            fig_fallback.add_trace(go.Scatter(
                x=[lon1, lon2], y=[lat1, lat2], mode="lines",
                line=dict(width=3, color="#dc2626" if speed_kmh > 150 else "#2563eb"),
                name=f"Distance ({dist_km:.1f} km)"
            ))
            fig_fallback.update_traces(textposition="top center")
            fig_fallback.update_xaxes(range=[min(lon1, lon2) - pad_lon, max(lon1, lon2) + pad_lon])
            fig_fallback.update_yaxes(range=[min(lat1, lat2) - pad_lat, max(lat1, lat2) + pad_lat])
            fig_fallback.update_layout(template="plotly_white")
            st.plotly_chart(fig_fallback, use_container_width=True)

        cm1, cm2, cm3 = st.columns(3)
        cm1.metric("Distance Between Points", f"{dist_km:.1f} km")
        cm2.metric("Recorded Time Gap", f"{int(time_gap_mins)} Minutes")
        if speed_kmh > 150:
            cm3.metric("Calculated Speed", f"{speed_kmh:.1f} km/h", delta="Physical Impossibility", delta_color="inverse")
        else:
            cm3.metric("Calculated Speed", f"{speed_kmh:.1f} km/h", delta="Legally Plausible", delta_color="normal")
            
    elif len(events_map_df) == 1:
        ev1 = events_map_df.iloc[0]
        st.info(f"Single location event recorded for `{selected_bundle}`: **{ev1['loc']}** at `{ev1['ts']}`")
    else:
        st.warning(f"No GIS coordinate data recorded for active bundle `{selected_bundle}`.")


# ----------------------------------------------------
# TAB 4: CHEAT SHEET
# ----------------------------------------------------
with tab_cheatsheet:
    st.subheader(f"Cross-Examination Cheat Sheet (`{selected_bundle}`)")
    if st.button("Generate Cheat Sheet Markdown", use_container_width=True):
        filepath = generate_cheatsheet(selected_bundle)
        st.success(f"Generated Cheat Sheet at: `{filepath}`")
        
    cs_file = f"CrossExam_CheatSheet_{selected_bundle}.md"
    if os.path.exists(cs_file):
        with open(cs_file, "r", encoding="utf-8") as f:
            st.markdown(f.read())
    else:
        st.warning(f"No Cheat Sheet generated yet for `{selected_bundle}`. Click the button above to generate!")
